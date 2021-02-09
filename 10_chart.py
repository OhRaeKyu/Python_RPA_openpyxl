from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

wb = load_workbook("sample.xlsx")
ws = wb.active

# bar_value = Reference(ws, min_row = 2, max_row = 11, min_col = 2, max_col = 3)  # B2 : C11 데이터를 차트 데이터로 생성
# bar_chart = BarChart()  # 차트 종류 설정
# bar_chart.add_data(bar_value) # 차트에 데이터 추가
# ws.add_chart(bar_chart, "E1") # 어떠한 차트를 어느 위치에 넣을지 정의

line_value = Reference(ws, min_row = 1, max_row = 11, min_col = 2, max_col = 3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True)  # 차트의 계열 이름
line_chart.title = "성적표" # 차트의 제목 
line_chart.style = 48 # 차트의 스타일
line_chart.y_axis.title = "점수"  # Y축의 제목
line_chart.x_axis.title = "번호"  # X축의 제목
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")